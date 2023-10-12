import openpyxl.chart.data_source
from openpyxl.chart import ScatterChart, Reference, Series
from openpyxl.chart.series_factory import SeriesFactory
from openpyxl.chart.series import SeriesLabel, StrRef
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import Font, CharacterProperties, Paragraph, ParagraphProperties
from openpyxl.utils import get_column_letter


def create_chart(ws, dataset_num):
    # initialize a chart object, set chart dimensions, set font and font size
    chart = ScatterChart()
    chart.height = 20.96
    chart.width = 31.44
    chart.graphical_properties = GraphicalProperties(ln=LineProperties(noFill=True))
    font = Font(typeface="Calibri")
    axis_font = CharacterProperties(latin=font, sz=2000)
    axis_label_font = CharacterProperties(latin=font, sz=1800)
    legend_font = CharacterProperties(latin=font, sz=1800)

    # legend properties
    chart.legend.overlay = True
    chart.legend.txPr = RichText(
        p=[Paragraph(pPr=ParagraphProperties(defRPr=legend_font), endParaRPr=legend_font)])

    # y-axis properties
    chart.y_axis.majorGridlines = None
    chart.y_axis.majorTickMark = "in"
    chart.y_axis.title = "mAU"
    chart.y_axis.title.tx.rich.p[0].pPr.defRPr = axis_font
    chart.y_axis.scaling.min = 0.0
    chart.y_axis.txPr = RichText(
        p=[Paragraph(pPr=ParagraphProperties(defRPr=axis_label_font), endParaRPr=axis_label_font)])
    chart.y_axis.number_format = "#0"

    # x-axis properties
    chart.x_axis.majorGridlines = None
    chart.x_axis.majorTickMark = "in"
    chart.x_axis.title = "Volume (mL)"
    chart.x_axis.title.tx.rich.p[0].pPr.defRPr = axis_font
    chart.x_axis.txPr = RichText(
        p=[Paragraph(pPr=ParagraphProperties(defRPr=axis_label_font), endParaRPr=axis_label_font)])
    chart.x_axis.number_format = "#0"
    chart.x_axis.scaling.min = 4.0
    chart.x_axis.scaling.max = 20.0
    chart.x_axis.majorUnit = 2.0

    # plot datasets
    for i in range(0, dataset_num):
        next_x = (i * 3) + 1
        x_values = Reference(ws, min_col=next_x, min_row=2, max_row=1000)
        next_y = next_x + 1
        y_values = Reference(ws, min_col=next_y, min_row=1, max_row=1000)
        series = SeriesFactory(y_values, x_values, title_from_data=True)
        # series.graphicalProperties.line.solidFill = "D22B2B"
        series.graphicalProperties.line.width = 60000
        chart.series.append(series)

    return chart


def get_cell_reference(ws, column_number):
    try:
        column_letter = get_column_letter(column_number)
        cell_reference = f"'{ws.title}'!${column_letter}$1"
        return cell_reference
    except Exception as e:
        print(f"Error while getting cell reference: {e}")
        return None
