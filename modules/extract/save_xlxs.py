import time

import openpyxl.chart.label
import openpyxl.utils.cell
import pandas as pd

from modules.extract.constants import *
from modules.extract.helpers import find_peak
from modules.extract.plotting import *


def generate_xlsx(input_filename, a260, a280, ratio):
    workbook = create_workbook(input_filename)
    writer = pd.ExcelWriter(workbook, engine="openpyxl", mode="a", if_sheet_exists="overlay")

    row = 0
    col = 0
    write_data(a260, A260_HEADER, writer, row, col)

    col += 3
    write_data(a280, A280_HEADER, writer, row, col)

    col += 3
    write_data(ratio, RATIO_HEADER, writer, row, col)

    peak_pos = find_peak(a260)
    chart_row = 2
    chart_col = chr(ord('A') + col + 3)
    worksheet = writer.sheets[WS_NAME]
    chart = create_chart(worksheet, len(a260), len(ratio), peak_pos, False)
    worksheet.add_chart(chart, f"{chart_col}{chart_row}")

    chart_row += 45
    chart = create_chart(worksheet, len(a260), len(ratio), peak_pos, True)
    worksheet.add_chart(chart, f"{chart_col}{chart_row}")

    writer.close()


def write_data(data, y_header, writer, row, col):
    df = create_dataframe(data, y_header)
    df.to_excel(writer, sheet_name=WS_NAME, index=False, startrow=row, startcol=col)


def create_dataframe(input_data, y_header):
    data = {VOLUME_HEADER: [], y_header: []}
    for x in input_data:
        data[VOLUME_HEADER].append(x)
        data[y_header].append(input_data[x])

    df = pd.DataFrame(data)

    return df


def create_workbook(filename):
    filename_suffix = time.strftime("%Y%m%d-%H%M%S")
    wb_filename = f"{filename[:-4]}_{filename_suffix}.xlsx"

    workbook = openpyxl.Workbook()
    workbook["Sheet"].title = WS_NAME

    workbook.save(wb_filename)

    return wb_filename
